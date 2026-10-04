# Office fixtures for twoody-server's document tests, written by programs
# other than Twoody's own (XlsxWriter, python-pptx), with fictional data.
#   python3 test/fixtures/office-fixtures.py test/fixtures   (XlsxWriter 3.2.9, python-pptx 1.0.2)
import datetime, io, re, sys, zipfile
import xlsxwriter
from pptx import Presentation
from pptx.util import Inches

out = sys.argv[1] if len(sys.argv) > 1 else "."

# ── budget-tessane.xlsx ─────────────────────────────────────────────────────
path = f"{out}/budget-tessane.xlsx"
wb = xlsxwriter.Workbook(path)
bold = wb.add_format({"bold": True})
title = wb.add_format({"bold": True, "align": "center"})
date = wb.add_format({"num_format": "dd/mm/yyyy"})
euro = wb.add_format({"num_format": '#,##0.00\\ "€"'})
percent = wb.add_format({"num_format": "0.0%"})
long_date = wb.add_format({"num_format": "[$-40C]d mmmm yyyy"})

sales = wb.add_worksheet("Ventes T1")
sales.merge_range("A1:F1", "Ventes du premier trimestre 2026 — Tessane", title)
for col, name in enumerate(["Date", "Client", "Produit", "Quantité", "Prix unitaire", "Montant"]):
    sales.write(1, col, name, bold)
rows = [
    (datetime.date(2026, 1, 5), "Claire Martin", "Abonnement annuel", 2, 120.0),
    (datetime.date(2026, 1, 19), "Atelier Lumière", "Formation", 1, 850.0),
    (datetime.date(2026, 2, 2), "Claire Martin", "Assistance", 3, 75.5),
    (datetime.date(2026, 2, 23), "Boulangerie Rivière", "Abonnement annuel", 1, 120.0),
    (datetime.date(2026, 3, 9), "Atelier Lumière", "Matériel", 4, 312.25),
    (datetime.date(2026, 3, 30), "Claire Martin", "Formation", 1, 850.0),
]
for i, (day, client, product, quantity, price) in enumerate(rows):
    r = i + 2
    sales.write_datetime(r, 0, datetime.datetime.combine(day, datetime.time()), date)
    sales.write(r, 1, client)
    sales.write(r, 2, product)
    sales.write_number(r, 3, quantity)
    sales.write_number(r, 4, price, euro)
    sales.write_formula(r, 5, f"=D{r+1}*E{r+1}", euro, quantity * price)
total = sum(q * p for _, _, _, q, p in rows)
sales.write(len(rows) + 2, 4, "Total", bold)
sales.write_formula(len(rows) + 2, 5, f"=SUM(F3:F{len(rows)+2})", euro, total)

summary = wb.add_worksheet("Synthèse")
summary.write_row(0, 0, ["Indicateur", 2025, 2026], bold)
summary.merge_range("A2:A3", "Chiffre d'affaires", None)
summary.write_row(1, 1, [9800, total], euro)
summary.write_row(2, 1, [0.12, 0.153], percent)
summary.write(3, 0, "Clôture de l'exercice")
summary.write_datetime(3, 2, datetime.datetime(2026, 12, 31), long_date)
summary.write(4, 0, "Objectif atteint")
summary.write_boolean(4, 2, True)
summary.write(5, 0, "Marge sur remise")
summary.write_formula(5, 2, "=1/0", None, "#DIV/0!")

notes = wb.add_worksheet("Notes")
notes.write(0, 0, "Remarques")
notes.write(1, 0, "Relancer Claire Martin pour le renouvellement.\nDevis envoyé le 2 mars.")
notes.write(2, 0, "Prix | remise négociée : voir le contrat")
notes.hide()
wb.close()

# Saved as Excel would save it: results stored, nothing left to compute.
data = io.BytesIO(open(path, "rb").read())
with zipfile.ZipFile(data) as source, zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as target:
    for item in source.infolist():
        body = source.read(item.filename)
        if item.filename == "xl/workbook.xml":
            body = re.sub(rb' fullCalcOnLoad="1"', b"", body)
        target.writestr(item, body)

# ── presentation-tessane.pptx ───────────────────────────────────────────────
deck = Presentation()
slide = deck.slides.add_slide(deck.slide_layouts[0])
slide.shapes.title.text = "Tessane — bilan du premier trimestre"
slide.placeholders[1].text = "Réunion du 2 avril 2026"

slide = deck.slides.add_slide(deck.slide_layouts[1])
slide.shapes.title.text = "Ce qui a marché"
body = slide.placeholders[1].text_frame
body.text = "Trois nouveaux clients, dont Atelier Lumière"
for line in ["Les formations : 1700 € sur le trimestre", "Claire Martin renouvelle son abonnement"]:
    body.add_paragraph().text = line
slide.notes_slide.notes_text_frame.text = "Dire que la formation de mars a été la plus rentable : marge de 41 %."

slide = deck.slides.add_slide(deck.slide_layouts[5])
slide.shapes.title.text = "Chiffres par client"
table = slide.shapes.add_table(3, 2, Inches(1), Inches(2), Inches(6), Inches(1.5)).table
for r, (a, b) in enumerate([("Client", "Montant"), ("Claire Martin", "1316.50 €"), ("Atelier Lumière", "2099.00 €")]):
    table.cell(r, 0).text = a
    table.cell(r, 1).text = b

slide = deck.slides.add_slide(deck.slide_layouts[6])  # blank: a picture-only slide, no text
slide = deck.slides.add_slide(deck.slide_layouts[1])
slide.shapes.title.text = "Prochaines étapes"
slide.placeholders[1].text = "Recruter un formateur avant juin"
slide.notes_slide.notes_text_frame.text = "Budget prévu : 3 000 € par mois. Ne pas annoncer le nom du candidat."
deck.save(f"{out}/presentation-tessane.pptx")
print("ok")
